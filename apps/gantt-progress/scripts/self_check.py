#!/usr/bin/env python3
"""Self-check for Gantt Progress v3 deliverables."""
from __future__ import annotations

import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
VBA = ROOT / "vba"
XLSX = ROOT / "GanttProgress_v3_TEMPLATE.xlsx"
errors: list[str] = []


def ok(msg: str) -> None:
    print(f"  OK  {msg}")


def fail(msg: str) -> None:
    errors.append(msg)
    print(f" FAIL {msg}")


def check_bas_files() -> None:
    required = {
        "modUtil.bas": ["SpeedOn", "SpeedOff", "CleanNames", "ReProtect"],
        "modValidate.bas": ["ValidateInput"],
        "modGantt.bas": ["DrawPlanLines"],
        "modProgress.bas": ["RecordStart", "RecordFinish", "UndoLast"],
        "modMain.bas": ["GenerateSheets"],
    }
    for fname, subs in required.items():
        path = VBA / fname
        if not path.exists():
            fail(f"missing {path}")
            continue
        text = path.read_text(encoding="utf-8")
        if "Option Explicit" not in text:
            fail(f"{fname}: missing Option Explicit")
        else:
            ok(f"{fname}: Option Explicit")
        for sub in subs:
            # Sub or Function
            if not re.search(rf"(Public |Private )?(Sub|Function) {sub}\b", text):
                fail(f"{fname}: missing procedure {sub}")
            else:
                ok(f"{fname}: has {sub}")


def check_cls() -> None:
    tw = VBA / "ThisWorkbook.cls"
    if not tw.exists():
        fail("missing ThisWorkbook.cls")
    else:
        t = tw.read_text(encoding="utf-8")
        if "Workbook_Open" not in t or "ReProtect" not in t:
            fail("ThisWorkbook.cls missing Workbook_Open/ReProtect")
        else:
            ok("ThisWorkbook.cls Workbook_Open → ReProtect")
    snip = VBA / "Worksheet_BeforeDoubleClick_snippet.bas"
    if not snip.exists():
        fail("missing Worksheet_BeforeDoubleClick_snippet.bas")
    else:
        t = snip.read_text(encoding="utf-8")
        if "Worksheet_BeforeDoubleClick" not in t or "Cancel" not in t:
            fail("snippet missing BeforeDoubleClick/Cancel")
        else:
            ok("BeforeDoubleClick snippet present")


def check_xlsx() -> None:
    try:
        from openpyxl import load_workbook
    except ImportError:
        fail("openpyxl not available")
        return
    if not XLSX.exists():
        fail(f"missing {XLSX}")
        return
    wb = load_workbook(XLSX)
    for name in ("Dashboard", "Input", "Template", "Log"):
        if name not in wb.sheetnames:
            fail(f"sheet missing: {name}")
        else:
            ok(f"sheet {name}")

    ws = wb["Template"]
    # Row 1 time headers from col 12 (L) — expect 288
    time_cols = 0
    for col in range(12, 12 + 300):
        v = ws.cell(1, col).value
        if v is None:
            break
        time_cols += 1
    if time_cols != 288:
        fail(f"Template time columns = {time_cols}, expected 288")
    else:
        ok("Template has 288 time columns (L+)")

    # Freeze
    if ws.freeze_panes != "L2":
        fail(f"Template freeze_panes={ws.freeze_panes}, expected L2")
    else:
        ok("Template freeze panes L2")

    if ws.sheet_state != "hidden":
        fail(f"Template not hidden ({ws.sheet_state})")
    else:
        ok("Template hidden")

    if wb["Log"].sheet_state != "hidden":
        fail("Log not hidden")
    else:
        ok("Log hidden")

    # Input headers at row 3
    ws_in = wb["Input"]
    expected = [
        "No", "日付", "大項目", "中項目", "小項目",
        "備考1", "備考2", "備考3", "備考4", "備考5",
        "開始時刻", "総時間(分)",
    ]
    for i, exp in enumerate(expected, 1):
        got = ws_in.cell(3, i).value
        if got != exp:
            fail(f"Input header col {i}: got {got!r}, expected {exp!r}")
        else:
            ok(f"Input header col {i}={exp}")

    # ListObject
    if not ws_in.tables:
        fail("Input has no ListObject table")
    else:
        ok(f"Input ListObject: {list(ws_in.tables.keys())}")

    # Sample rows
    data_rows = 0
    for r in range(4, 30):
        if ws_in.cell(r, 2).value is not None:
            data_rows += 1
    if data_rows < 10:
        fail(f"Input sample rows={data_rows}, expected >=10")
    else:
        ok(f"Input sample rows={data_rows}")

    # Named ranges on Template should be none (workbook-level check for Template refs)
    # openpyxl defined_names
    bad = []
    for dn in wb.defined_names.values():
        try:
            dest = dn.attr_text
        except Exception:
            dest = str(dn)
        if dest and "Template!" in str(dest):
            bad.append(dest)
    if bad:
        fail(f"Template has named ranges: {bad}")
    else:
        ok("No Template named ranges")


def main() -> int:
    print("=== Gantt Progress v3 self-check ===")
    check_bas_files()
    check_cls()
    check_xlsx()
    print("---")
    if errors:
        print(f"FAILED: {len(errors)} issue(s)")
        for e in errors:
            print(" -", e)
        return 1
    print("ALL CHECKS PASSED")
    return 0


if __name__ == "__main__":
    sys.exit(main())
