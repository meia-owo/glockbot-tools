#!/usr/bin/env python3
"""Build GanttProgress_v3_TEMPLATE.xlsx with openpyxl (no macros)."""
from __future__ import annotations

from datetime import datetime, time, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.chart.shapes import GraphicalProperties  # noqa: F401 — keep import soft

OUT = Path(__file__).resolve().parents[1] / "GanttProgress_v3_TEMPLATE.xlsx"

# Colors
FILL_HEADER = PatternFill("solid", fgColor="37474F")
FONT_HEADER = Font(color="FFFFFF", bold=True, size=10, name="Meiryo UI")
FILL_NOTICE = PatternFill("solid", fgColor="FFF9C4")
FILL_PLAN = PatternFill("solid", fgColor="C8C8C8")
THIN = Border(
    left=Side(style="thin", color="BDBDBD"),
    right=Side(style="thin", color="BDBDBD"),
    top=Side(style="thin", color="BDBDBD"),
    bottom=Side(style="thin", color="BDBDBD"),
)
THICK_RIGHT = Border(
    left=Side(style="thin", color="BDBDBD"),
    right=Side(style="medium", color="424242"),
    top=Side(style="thin", color="BDBDBD"),
    bottom=Side(style="thin", color="BDBDBD"),
)


def build() -> Path:
    wb = Workbook()

    # --- Dashboard ---
    ws_dash = wb.active
    ws_dash.title = "Dashboard"
    _build_dashboard(ws_dash)

    # --- Input ---
    ws_in = wb.create_sheet("Input")
    _build_input(ws_in)

    # --- Template (hidden later) ---
    ws_tmpl = wb.create_sheet("Template")
    _build_template(ws_tmpl)

    # --- Log (hidden later) ---
    ws_log = wb.create_sheet("Log")
    _build_log(ws_log)

    ws_tmpl.sheet_state = "hidden"
    ws_log.sheet_state = "hidden"

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    return OUT


def _build_dashboard(ws) -> None:
    ws["A1"] = "ガントチャート進捗管理ツール v3 — Dashboard"
    ws["A1"].font = Font(bold=True, size=16, name="Meiryo UI")
    ws.merge_cells("A1:F1")

    ws["A3"] = "■ 全体進捗"
    ws["A3"].font = Font(bold=True, size=12)
    ws["A4"] = "完了タスク数"
    ws["B4"] = "(日付シート生成後、下記の日別サマリ／数式で集計)"
    ws["A5"] = "全体タスク数"
    ws["A6"] = "進捗率"
    ws["B6"] = "※ M365: 各日付シート I列 COUNTIF の合算 / COUNTA で算出"
    ws["C6"] = "例: =完了数/全体数"

    ws["A8"] = "■ 遅延タスク一覧（計画終了超過かつ未完了）"
    ws["A8"].font = Font(bold=True, size=12)
    ws["A9"] = "※ 日付シート生成後、FILTER/COUNTIFS 等で自動抽出を推奨。"
    ws["A10"] = "※ タイムスタンプ記録の都度の全再集計は行わない（性能）。"

    ws["A12"] = "■ 日別サマリ（ハイパーリンク）"
    ws["A12"].font = Font(bold=True, size=12)
    ws["A13"] = "シート名"
    ws["B13"] = "完了率メモ"
    ws["C13"] = "備考"
    for col in ("A", "B", "C"):
        ws[f"{col}13"].fill = FILL_HEADER
        ws[f"{col}13"].font = FONT_HEADER

    ws["A15"] = "「更新」ボタン → RefreshDashboard（VBA）で日別リンク再構築"
    ws["A17"] = "操作ガイド"
    ws["A17"].font = Font(bold=True)
    ws["A18"] = "1. Input シートにタスクを記入（大項目レーンごと時系列順）"
    ws["A19"] = "2. 「シート生成」ボタンで YYYYMMDD シートを生成"
    ws["A20"] = "3. 日付シートの J/K 列をダブルクリックで着工/完了を記録"
    ws["A21"] = "4. 本 Dashboard で全体進捗を確認"

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["C"].width = 30
    ws.column_dimensions["D"].width = 18
    ws.freeze_panes = "A3"


def _build_input(ws) -> None:
    # Row 1: notice
    ws["A1"] = (
        "【注意】タスクは大項目（レーン）ごとに時系列順（上から実行順）で記入してください。"
        "異なる大項目の行が交互に並んでも構いません（チェーン計算は大項目ごとに独立）。"
        "各日付×各大項目の先頭タスクには開始時刻（アンカー）が必須です。"
        "備考1〜5の見出しは自由にリネーム可（VBAは列位置F〜Jで参照）。"
    )
    ws["A1"].fill = FILL_NOTICE
    ws["A1"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.merge_cells("A1:L1")
    ws.row_dimensions[1].height = 48

    # Row 2: blank spacer for button area note
    ws["A2"] = "→ 「シート生成」ボタンをここに配置し、GenerateSheets を割り当ててください。"
    ws.merge_cells("A2:L2")

    # Row 3: headers (table starts here — ListObject)
    headers = [
        "No",
        "日付",
        "大項目",
        "中項目",
        "小項目",
        "備考1",
        "備考2",
        "備考3",
        "備考4",
        "備考5",
        "開始時刻",
        "総時間(分)",
    ]
    for i, h in enumerate(headers, 1):
        cell = ws.cell(3, i, h)
        cell.fill = FILL_HEADER
        cell.font = FONT_HEADER
        cell.alignment = Alignment(horizontal="center")

    # Sample data: 2 dates, 2 lanes, multi-anchor
    # Lane ラインA: morning chain + afternoon re-anchor at 13:00
    # Lane ラインB: parallel
    samples = [
        # date, large, mid, small, n1,n2,n3,n4,n5, start, mins
        (datetime(2026, 7, 1), "ラインA", "段取り", "治具セット", "担当:田中", "設備-01", "", "", "", time(9, 0), 30),
        (datetime(2026, 7, 1), "ラインA", "段取り", "材料搬入", "", "", "", "", "", None, 20),
        (datetime(2026, 7, 1), "ラインA", "加工", "粗削り", "注意:冷却", "", "", "", "", None, 45),
        (datetime(2026, 7, 1), "ラインA", "加工", "仕上げ", "", "", "", "", "", None, 40),
        # re-anchor after lunch
        (datetime(2026, 7, 1), "ラインA", "検査", "寸法測定", "午後工程", "", "", "", "", time(13, 0), 25),
        (datetime(2026, 7, 1), "ラインA", "検査", "外観検査", "", "", "", "", "", None, 15),
        # Lane B parallel
        (datetime(2026, 7, 1), "ラインB", "準備", "プログラム確認", "担当:鈴木", "CNC-2", "", "", "", time(8, 30), 35),
        (datetime(2026, 7, 1), "ラインB", "準備", "ツールセット", "", "", "", "", "", None, 25),
        (datetime(2026, 7, 1), "ラインB", "加工", "メイン加工", "", "", "", "", "", None, 90),
        (datetime(2026, 7, 1), "ラインB", "後工程", "バリ取り", "", "", "", "", "", time(14, 0), 30),
        # Day 2
        (datetime(2026, 7, 2), "ラインA", "組立", "部品組付", "担当:田中", "", "", "", "", time(9, 0), 60),
        (datetime(2026, 7, 2), "ラインA", "組立", "配線", "", "", "", "", "", None, 45),
        (datetime(2026, 7, 2), "ラインB", "検査", "最終検査", "担当:佐藤", "", "", "", "", time(10, 0), 50),
        (datetime(2026, 7, 2), "ラインB", "出荷", "梱包", "", "", "", "", "", None, 30),
    ]

    for idx, row in enumerate(samples, 1):
        r = 3 + idx  # data starts row 4
        ws.cell(r, 1, idx)
        ws.cell(r, 2, row[0]).number_format = "yyyy/mm/dd"
        ws.cell(r, 3, row[1])
        ws.cell(r, 4, row[2])
        ws.cell(r, 5, row[3])
        for j in range(5):
            ws.cell(r, 6 + j, row[4 + j] or None)
        if row[9] is not None:
            ws.cell(r, 11, row[9]).number_format = "h:mm"
        ws.cell(r, 12, row[10])

    last_data = 3 + len(samples)
    # ListObject table
    table = Table(displayName="InputTable", ref=f"A3:L{last_data}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False,
        showLastColumn=False, showRowStripes=True, showColumnStripes=False
    )
    ws.add_table(table)

    # Data validation: 5-min start times (as list of times — use a helper range on far columns)
    # Build time list on columns AA (hidden-ish)
    for i in range(288):
        mins = i * 5
        t = time(mins // 60, mins % 60)
        ws.cell(1 + i, 27, t).number_format = "h:mm"  # AA
    ws.column_dimensions["AA"].hidden = True
    dv = DataValidation(
        type="list",
        formula1="=$AA$1:$AA$288",
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="開始時刻",
        error="5分刻みの時刻を選択してください。",
    )
    dv.add(f"K4:K{max(last_data, 100)}")
    ws.add_data_validation(dv)

    widths = {
        "A": 6, "B": 12, "C": 12, "D": 12, "E": 14,
        "F": 14, "G": 12, "H": 10, "I": 10, "J": 10,
        "K": 10, "L": 12,
    }
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    ws.freeze_panes = "A4"


def _build_template(ws) -> None:
    """Date-sheet layout: A-C hierarchy, D-H notes, I status, J start, K finish, L+ 288 times."""
    headers_left = [
        "大項目", "中項目", "小項目",
        "備考1", "備考2", "備考3", "備考4", "備考5",
        "ステータス", "着工", "完了",
    ]
    for i, h in enumerate(headers_left, 1):
        cell = ws.cell(1, i, h)
        cell.fill = FILL_HEADER
        cell.font = FONT_HEADER
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    # Time headers 0:00 .. 23:55
    for i in range(288):
        mins = i * 5
        col = 12 + i  # L = 12
        tval = time(mins // 60, mins % 60)
        cell = ws.cell(1, col, tval)
        cell.number_format = "h:mm"
        cell.fill = FILL_HEADER
        cell.font = Font(color="FFFFFF", bold=True, size=8, name="Meiryo UI")
        cell.alignment = Alignment(horizontal="center", textRotation=90)
        # Hourly thick border on the last 5-min of each hour (xx:55) right edge,
        # or on xx:00 left — use medium right border every 12th col (end of hour)
        if (i + 1) % 12 == 0:
            cell.border = THICK_RIGHT
        else:
            cell.border = THIN
        ws.column_dimensions[get_column_letter(col)].width = 2.5

    # Sample empty data row styling (row 2 placeholder)
    for c in range(1, 12):
        ws.cell(2, c).border = THIN

    # Column widths left side
    for col, w in zip("ABCDEFGHIJK", [12, 12, 14, 10, 10, 10, 10, 10, 10, 18, 18]):
        ws.column_dimensions[col].width = w

    # Freeze panes: columns A-K + header row
    ws.freeze_panes = "L2"

    # Row height for header
    ws.row_dimensions[1].height = 40

    # Button placement notes (row 0 area — put in A2 comment area as text above data)
    # Use a note in cell far below or in sheet properties — put instruction in row 1 merged note via comment
    from openpyxl.comments import Comment
    ws["A1"].comment = Comment(
        "フォームコントロール配置例:\n"
        "着工→ButtonRecordStart / 完了→ButtonRecordFinish / "
        "記録取り消し→ButtonUndoLast / 階層1〜3→ShowLevel1〜3\n"
        "ダブルクリック: J=着工, K=完了（シートモジュールにイベント貼付）",
        "Gantt v3",
    )

    # Group note columns D-H outline (columns group) so they can be collapsed
    ws.column_dimensions.group("D", "H", hidden=False)

    # NO named ranges on Template (spec)


def _build_log(ws) -> None:
    headers = ["日時", "シート名", "行", "操作種別", "タスク"]
    for i, h in enumerate(headers, 1):
        cell = ws.cell(1, i, h)
        cell.fill = FILL_HEADER
        cell.font = FONT_HEADER
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 8
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 40
    ws.freeze_panes = "A2"


if __name__ == "__main__":
    path = build()
    print(f"Wrote {path} ({path.stat().st_size} bytes)")
